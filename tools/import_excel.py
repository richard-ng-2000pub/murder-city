#!/usr/bin/env python3
"""
Convert your Murder City Excel map to data/cell-legends.js for GitHub Pages.

Expected information per row:
  district number, avenue letter, street number, and legend image number/file.

The importer accepts many common column names automatically. Examples:
  District | Avenue | Street | Legend
  district_no | avenue | street | legend_id
  District | Column | Row | Image

Legend values can be: 1, 01, legend-01, legend-01.png, etc.
"""
from __future__ import annotations
import argparse
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Missing dependency: openpyxl. Run: pip install -r tools/requirements.txt", file=sys.stderr)
    raise SystemExit(2)

ALIASES = {
    "district": {"district", "districtno", "districtnumber", "districtid", "zone", "area"},
    "avenue": {"avenue", "ave", "column", "col", "avenueletter"},
    "street": {"street", "streetno", "streetnumber", "row", "rownumber"},
    "legend": {"legend", "legendid", "legendno", "legendnumber", "image", "imagefile", "filename", "png", "icon", "iconid"},
}

def canon(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())

def find_columns(headers):
    found = {}
    canon_headers = [canon(h) for h in headers]
    for field, aliases in ALIASES.items():
        for idx, h in enumerate(canon_headers):
            if h in aliases:
                found[field] = idx
                break
    return found

def legend_num(value) -> int | None:
    if value is None: return None
    m = re.search(r"(\d{1,2})", str(value))
    if not m: return None
    n = int(m.group(1))
    return n if 1 <= n <= 35 else None

def normalize_key(district, avenue, street):
    try:
        d = str(int(float(str(district).strip())))
        s = str(int(float(str(street).strip())))
    except Exception:
        return None
    a = re.sub(r"[^A-Za-z]", "", str(avenue or "")).upper()
    if not a: return None
    return f"{d}|{a}|{s}"

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", help="Path to the city-map Excel workbook")
    ap.add_argument("--sheet", help="Worksheet name. Defaults to active sheet.")
    ap.add_argument("--output", default="data/cell-legends.js")
    args = ap.parse_args()

    wb = load_workbook(args.xlsx, data_only=True, read_only=True)
    ws = wb[args.sheet] if args.sheet else wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        headers = list(next(rows))
    except StopIteration:
        raise SystemExit("Workbook is empty.")

    cols = find_columns(headers)
    missing = [x for x in ("district","avenue","street","legend") if x not in cols]
    if missing:
        print("Could not detect columns:", ", ".join(missing), file=sys.stderr)
        print("Headers found:", headers, file=sys.stderr)
        print("Rename columns to District, Avenue, Street, Legend — or send the Excel file and the importer can be adapted to its exact layout.", file=sys.stderr)
        raise SystemExit(3)

    mapping = {}
    skipped = 0
    duplicates = 0
    for row in rows:
        key = normalize_key(row[cols["district"]], row[cols["avenue"]], row[cols["street"]])
        lid = legend_num(row[cols["legend"]])
        if not key or lid is None:
            skipped += 1
            continue
        if key in mapping and mapping[key] != lid:
            duplicates += 1
        mapping[key] = lid

    out = Path(args.output)
    if not out.is_absolute():
        project_root = Path(__file__).resolve().parents[1]
        out = project_root / out
    out.parent.mkdir(parents=True, exist_ok=True)
    payload = "window.MURDER_CITY_CELL_LEGENDS = " + json.dumps({"mode":"exact","cells":mapping}, indent=2, sort_keys=True) + ";\n"
    out.write_text(payload, encoding="utf-8")

    print(f"Wrote {len(mapping):,} exact cell mappings -> {out}")
    if skipped: print(f"Skipped {skipped:,} incomplete/unrecognized rows.")
    if duplicates: print(f"Warning: {duplicates:,} duplicate addresses had conflicting legend IDs; the last row won.")

if __name__ == "__main__":
    main()
